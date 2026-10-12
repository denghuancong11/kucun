"""Independently inspect runtime downloads, optionally comparing the supplied original XLSX."""
import datetime, hashlib, json, sys, zipfile
import openpyxl

file, reference, purpose = sys.argv[1:]
w = openpyxl.load_workbook(file)
assert len(w.sheetnames) == 1
s = w.active
is_move = s.title == '移仓升级'
assert is_move or s.title == '转仓升级'
headers = ['型号', '来源', '发货计划号', '发货时间', '版本号', '已贴FNSKU'] + (
    ['初始调拨数量', 'FBA剩余库存', '调拨店铺', '状况', 'RMA', '移仓地址', '订单号', '承运商', '运单号', '移仓-已发货数量', '升级中数量', '升级完数量', '升级完，版本号'] if is_move else
    ['退仓数量', '目前所在海外仓', '退仓店铺', '状况', 'RMA', '实际清点数量', '升级中数量', '升级完数量', '升级完，版本号'])
start = 3 if is_move else 4
assert s.max_column == len(headers)
assert [c.value for c in s[2]] == headers
assert all(c.value is None for c in s[1])
merges = sorted(map(str, s.merged_cells.ranges))
assert merges == ([] if is_move else sorted(f'{chr(65+i)}2:{chr(65+i)}3' for i in range(15)))
assert not w.defined_names and not s.protection.sheet
assert all(not d.hidden for d in s.column_dimensions.values())
assert all(not d.hidden for d in s.row_dimensions.values())
for row in s:
    for cell in row:
        assert cell.comment is None and cell.data_type != 'f'
        assert not any(t in str(cell.value) for t in ['XXX', '填写项', '记录版本', '移仓单号', '转仓单号'])
assert all(c.value is None for c in s[3]) if not is_move else True
properties = {p.name: p.value for p in w.custom_doc_props}
assert properties == ({'AsterWorkbookPurpose': 'relocation-waybills-view-only'} if purpose == 'waybill' else {})
if reference:
    assert hashlib.sha256(open(reference, 'rb').read()).hexdigest() == '0483223d99e127e7f52b0f37690b40bca0f7a028bdb28fac781361e0b326d09d'
    original = openpyxl.load_workbook(reference)[s.title]
    assert s.sheet_format == original.sheet_format
    assert s.sheet_view.zoomScale == original.sheet_view.zoomScale
    assert s.page_margins == original.page_margins and s.page_setup == original.page_setup
    assert {k: dict(v) for k, v in s.column_dimensions.items()} == {k: dict(v) for k, v in original.column_dimensions.items()}
    for row in range(2, start):
        assert s.row_dimensions[row].height == original.row_dimensions[row].height
        for col in range(1, len(headers)+1):
            assert s.cell(row, col)._style == original.cell(row, col)._style
    for row in range(start, s.max_row+1):
        assert s.row_dimensions[row].height == original.row_dimensions[start].height
        for col in range(1, len(headers)+1):
            assert s.cell(row, col)._style == original.cell(start, col)._style
    with zipfile.ZipFile(reference) as a, zipfile.ZipFile(file) as b:
        assert a.read('xl/styles.xml') == b.read('xl/styles.xml')
        assert a.read('xl/theme/theme1.xml') == b.read('xl/theme/theme1.xml')
print(json.dumps({'sheet': s.title, 'columns': s.max_column, 'dataStartRow': start, 'merges': merges,
                  'rows': [[c.value for c in row] for row in s.iter_rows(min_row=start)] if s.max_row >= start else [],
                  'comparedWithOriginal': bool(reference), 'purpose': purpose}, ensure_ascii=False, default=str))
