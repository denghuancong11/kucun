"""Independent read-only verification of actual exported files (openpyxl + ZIP CRC)."""
import json
import sys
import zipfile
from openpyxl import load_workbook

with zipfile.ZipFile(sys.argv[1]) as archive:
    assert archive.testzip() is None
book = load_workbook(sys.argv[1], data_only=False)
sheet = book.active
result = {
    'sheets': book.sheetnames,
    'maxRow': sheet.max_row,
    'maxColumn': sheet.max_column,
    'rows': [[cell.value for cell in row] for row in sheet.iter_rows()],
    'types': [[cell.data_type for cell in row] for row in sheet.iter_rows()],
    'numberFormats': [[cell.number_format for cell in row] for row in sheet.iter_rows()],
    'freezePanes': sheet.freeze_panes,
}
print(json.dumps(result, ensure_ascii=True))
book.close()